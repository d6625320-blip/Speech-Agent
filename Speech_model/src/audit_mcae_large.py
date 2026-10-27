"""审计 MCAE 数据、标签、音频质量，并生成可复现的 LOSO 划分。"""

from __future__ import annotations

import argparse
import csv
import json
import math
import random
import wave
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import numpy as np
import yaml
from openpyxl import load_workbook

from config_utils import load_config, resolve_project_path


LABELS = ["angry", "disgust", "fear", "happy", "neutral", "sad", "surprised"]
SOURCE_TO_LABEL = {
    "anger": "angry",
    "disgust": "disgust",
    "fear": "fear",
    "happiness": "happy",
    "neutral": "neutral",
    "sadness": "sad",
    "surprise": "surprised",
}
LABEL_TO_FOLDER = {
    "angry": "Anger",
    "disgust": "Disgust",
    "fear": "Fear",
    "happy": "Happiness",
    "neutral": "Neutral",
    "sad": "Sadness",
    "surprised": "Surprise",
}
MODEL_OUTPUT_ORDER = [
    "angry",
    "disgusted",
    "fearful",
    "happy",
    "neutral",
    "other",
    "sad",
    "surprised",
    "unknown",
]


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)


def excel_rows(path: Path, sheet: str | None = None) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    rows = [dict(zip(headers, values)) for values in iterator if any(value is not None for value in values)]
    workbook.close()
    return rows


def sample_id(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def number(value: Any) -> float:
    if value is None or str(value).strip() == "":
        return math.nan
    return float(value)


def percentile_summary(values: list[float]) -> dict[str, float]:
    array = np.asarray([value for value in values if math.isfinite(value)], dtype=np.float64)
    if array.size == 0:
        return {}
    return {
        "min": float(array.min()),
        "p01": float(np.percentile(array, 1)),
        "p05": float(np.percentile(array, 5)),
        "median": float(np.median(array)),
        "mean": float(array.mean()),
        "p95": float(np.percentile(array, 95)),
        "p99": float(np.percentile(array, 99)),
        "max": float(array.max()),
    }


def decode_pcm(path: Path) -> tuple[np.ndarray, dict[str, Any]]:
    with wave.open(str(path), "rb") as audio:
        channels = audio.getnchannels()
        sample_rate = audio.getframerate()
        sample_width = audio.getsampwidth()
        frames = audio.getnframes()
        compression = audio.getcomptype()
        raw = audio.readframes(frames)
    if compression != "NONE":
        raise ValueError(f"不支持压缩 WAV: {path} ({compression})")
    if sample_width == 1:
        values = (np.frombuffer(raw, dtype=np.uint8).astype(np.float32) - 128.0) / 128.0
    elif sample_width == 2:
        values = np.frombuffer(raw, dtype="<i2").astype(np.float32) / 32768.0
    elif sample_width == 4:
        values = np.frombuffer(raw, dtype="<i4").astype(np.float32) / 2147483648.0
    else:
        raise ValueError(f"不支持 {sample_width * 8}-bit WAV: {path}")
    if frames <= 0 or values.size != frames * channels:
        raise ValueError(f"WAV 帧数无效: {path}")
    waveform = values.reshape(frames, channels)
    return waveform, {
        "channels": channels,
        "sample_rate": sample_rate,
        "sample_width_bytes": sample_width,
        "frames": frames,
        "compression": compression,
    }


def consecutive_silence_ratio(mask: np.ndarray, from_start: bool) -> float:
    sequence = mask if from_start else mask[::-1]
    non_silent = np.flatnonzero(~sequence)
    count = int(non_silent[0]) if non_silent.size else len(sequence)
    return count / max(1, len(sequence))


def actor_gender_map(rater_path: Path) -> dict[int, str]:
    rows = excel_rows(rater_path, "Sheet1")
    values: dict[int, set[str]] = defaultdict(set)
    for row in rows:
        if row.get("actress") is not None and row.get("actresssex") is not None:
            values[int(row["actress"])].add(str(row["actresssex"]).strip().lower())
    ambiguous = {actor: sexes for actor, sexes in values.items() if len(sexes) != 1}
    if ambiguous:
        raise ValueError(f"演员性别 metadata 不一致: {ambiguous}")
    return {actor: next(iter(sexes)) for actor, sexes in values.items()}


def build_master_rows(
    metadata_path: Path,
    audio_dir: Path,
    manifest_path: Path,
    gender_by_actor: dict[int, str],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    metadata = excel_rows(metadata_path)
    with manifest_path.open(newline="", encoding="utf-8-sig") as handle:
        manifest = list(csv.DictReader(handle))
    manifest_by_id = {row["sample_id"]: row for row in manifest}
    if len(manifest_by_id) != len(manifest):
        raise ValueError("现有 manifest 存在重复 sample_id")

    master: list[dict[str, Any]] = []
    mismatch_rows: list[dict[str, Any]] = []
    vote_mismatches: list[dict[str, Any]] = []
    vote_ties = 0
    vote_totals: Counter[int] = Counter()
    recognition_rate_mismatches = 0
    seen: set[str] = set()
    for source in metadata:
        sid = sample_id(source["Filename"])
        if sid in seen:
            raise ValueError(f"原始 metadata 存在重复 Filename: {sid}")
        seen.add(sid)
        source_label = str(source["emotion"]).strip()
        label = SOURCE_TO_LABEL.get(source_label.lower())
        if label is None:
            raise ValueError(f"未知原始标签: {source_label}")
        actor = int(source["Actress"])
        vote_names = ["Neutral", "Disgust", "Fear", "Sadness", "Surprise", "Anger", "Happiness"]
        votes = {name: int(source[name]) for name in vote_names}
        vote_total = sum(votes.values())
        vote_totals[vote_total] += 1
        maximum_vote = max(votes.values())
        winners = [name for name, value in votes.items() if value == maximum_vote]
        vote_ties += int(len(winners) > 1)
        if source_label.lower() not in {winner.lower() for winner in winners}:
            vote_mismatches.append({"sample_id": sid, "source_label": source_label, "vote_winners": winners, "votes": votes})
        recognition_rate = number(source.get("Recognition rate"))
        if vote_total <= 0 or abs(recognition_rate - votes[source_label.title()] / vote_total) > 1e-9:
            recognition_rate_mismatches += 1
        expected_path = audio_dir / LABEL_TO_FOLDER[label] / f"{sid}.wav"
        manifest_row = manifest_by_id.get(sid)
        checks = {
            "metadata_audio_exists": expected_path.is_file(),
            "manifest_exists": manifest_row is not None,
            "manifest_label_match": manifest_row is not None and manifest_row["label"] == label,
            "manifest_actor_match": manifest_row is not None and int(manifest_row["actor_id"]) == actor,
            "folder_label_match": expected_path.parent.name == LABEL_TO_FOLDER[label],
        }
        if not all(checks.values()):
            mismatch_rows.append({"sample_id": sid, "label": label, "actor_id": actor, **checks})

        max_f0 = number(source.get("max_f0"))
        min_f0 = number(source.get("min_f0"))
        row = {
            "sample_id": sid,
            "audio_path": str(expected_path.resolve()),
            "label": label,
            "label_id": LABELS.index(label),
            "source_label": source_label,
            "actor_id": actor,
            "speaker_id": actor,
            "loso_test_fold": actor,
            "actor_gender": gender_by_actor.get(actor, "unknown"),
            "intensity": number(source.get("Intensity")),
            "recognition_rate": recognition_rate,
            "sentence_id": sample_id(source.get("sentencecode")),
            "transcript": str(source.get("sentence content") or "").strip(),
            "sentence_group": sample_id(source.get("groupnumber")),
            "subject_name": str(source.get("subjectname") or "").strip(),
            "metadata_duration": number(source.get("duration")),
            "f0_mean": number(source.get("meanF0")),
            "f0_range": max_f0 - min_f0,
            "hnr": number(source.get("hnr")),
            "jitter": number(source.get("localJitter")),
            "shimmer": number(source.get("localShimmer")),
            "spectral_cog": number(source.get("spectralCOG")),
            "spectral_spread": number(source.get("spectralSpread")),
        }
        master.append(row)

    extra_manifest = sorted(set(manifest_by_id) - seen)
    audit = {
        "metadata_rows": len(metadata),
        "manifest_rows": len(manifest),
        "master_rows": len(master),
        "metadata_unique_ids": len(seen),
        "manifest_unique_ids": len(manifest_by_id),
        "metadata_manifest_mismatches": len(mismatch_rows),
        "extra_manifest_ids": extra_manifest,
        "mismatch_examples": mismatch_rows[:20],
        "vote_label_mismatches": len(vote_mismatches),
        "vote_ties": vote_ties,
        "vote_totals": dict(vote_totals),
        "recognition_rate_vote_mismatches": recognition_rate_mismatches,
        "label_definition": "emotion equals the unique plurality label among 40 human votes; recognition_rate equals votes_for_label / 40",
    }
    return master, audit


def audit_audio(
    rows: list[dict[str, Any]], silence_dbfs: float, clipping_threshold: float
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    silence_threshold = 10.0 ** (silence_dbfs / 20.0)
    details: list[dict[str, Any]] = []
    format_counts: Counter[str] = Counter()
    errors: list[dict[str, str]] = []
    for index, source in enumerate(rows, start=1):
        path = Path(source["audio_path"])
        try:
            waveform, header = decode_pcm(path)
        except Exception as exc:
            errors.append({"sample_id": source["sample_id"], "audio_path": str(path), "error": str(exc)})
            continue
        mono = waveform.mean(axis=1, dtype=np.float32)
        absolute = np.abs(mono)
        silent = absolute <= silence_threshold
        channel_difference = 0.0
        channel_correlation = 1.0
        if header["channels"] == 2:
            channel_difference = float(np.max(np.abs(waveform[:, 0] - waveform[:, 1])))
            left = waveform[:, 0].astype(np.float64)
            right = waveform[:, 1].astype(np.float64)
            if left.std() > 0 and right.std() > 0:
                channel_correlation = float(np.corrcoef(left, right)[0, 1])
        duration = header["frames"] / header["sample_rate"]
        format_key = (
            f"channels={header['channels']},sample_rate={header['sample_rate']},"
            f"bits={header['sample_width_bytes'] * 8},compression={header['compression']}"
        )
        format_counts[format_key] += 1
        details.append(
            {
                "sample_id": source["sample_id"],
                "label": source["label"],
                "actor_id": source["actor_id"],
                "audio_path": str(path),
                **header,
                "duration_sec": duration,
                "metadata_duration_sec": source["metadata_duration"],
                "duration_difference_sec": duration - source["metadata_duration"],
                "mono_float32_samples": len(mono),
                "expected_16k_samples": int(round(duration * 16000)),
                "wrong_flatten_duration_sec": waveform.size / header["sample_rate"],
                "peak_abs": float(absolute.max()),
                "rms": float(np.sqrt(np.mean(np.square(mono, dtype=np.float64)))),
                "clipping_ratio": float(np.mean(absolute >= clipping_threshold)),
                "silence_ratio": float(silent.mean()),
                "leading_silence_ratio": consecutive_silence_ratio(silent, True),
                "trailing_silence_ratio": consecutive_silence_ratio(silent, False),
                "channel_max_abs_difference": channel_difference,
                "channel_correlation": channel_correlation,
            }
        )
        if index % 1000 == 0:
            print(f"audio_audit={index}/{len(rows)}")

    numeric_fields = [
        "duration_sec",
        "duration_difference_sec",
        "peak_abs",
        "rms",
        "clipping_ratio",
        "silence_ratio",
        "leading_silence_ratio",
        "trailing_silence_ratio",
        "channel_max_abs_difference",
        "channel_correlation",
    ]
    summary = {
        "files_checked": len(details),
        "errors": errors,
        "format_counts": dict(format_counts),
        "float_conversion": "PCM converted to float32 in [-1, 1]",
        "correct_channel_conversion": "mean over channel axis, output shape [time]",
        "wrong_flatten_effect": "stereo flatten would double sample count and interleave channels",
        "target_sample_rate": 16000,
        "vad_enabled": False,
        "silence_threshold_dbfs": silence_dbfs,
        "statistics": {
            field: percentile_summary([float(row[field]) for row in details])
            for field in numeric_fields
        },
        "by_label": {
            label: {
                "samples": sum(row["label"] == label for row in details),
                **{
                    field: percentile_summary(
                        [float(row[field]) for row in details if row["label"] == label]
                    )
                    for field in numeric_fields
                },
            }
            for label in LABELS
        },
    }
    return details, summary


def export_listening_samples(
    selected: list[dict[str, Any]], output_dir: Path, target_sample_rate: int
) -> list[dict[str, Any]]:
    try:
        import torch
        import torchaudio
    except ImportError:
        torch = None
        torchaudio = None

    exported: list[dict[str, Any]] = []
    for source in selected:
        target = output_dir / "listening_samples" / source["label"] / f"{source['sample_id']}.wav"
        target.parent.mkdir(parents=True, exist_ok=True)
        if torchaudio is not None and torch is not None:
            waveform, sample_rate = torchaudio.load(source["audio_path"])
            if waveform.ndim != 2:
                raise ValueError(f"音频张量维度异常: {source['audio_path']}")
            waveform = waveform.to(torch.float32).mean(dim=0, keepdim=True)
            if sample_rate != target_sample_rate:
                waveform = torchaudio.functional.resample(waveform, sample_rate, target_sample_rate)
            torchaudio.save(str(target), waveform.cpu(), target_sample_rate, encoding="PCM_S", bits_per_sample=16)
        else:
            from math import gcd
            from scipy.signal import resample_poly

            samples, header = decode_pcm(Path(source["audio_path"]))
            mono = samples.mean(axis=1, dtype=np.float32)
            common = gcd(int(header["sample_rate"]), target_sample_rate)
            mono = resample_poly(mono, target_sample_rate // common, int(header["sample_rate"]) // common).astype(np.float32)
            pcm = np.round(np.clip(mono, -1.0, 1.0) * 32767.0).astype("<i2")
            with wave.open(str(target), "wb") as audio:
                audio.setnchannels(1)
                audio.setsampwidth(2)
                audio.setframerate(target_sample_rate)
                audio.writeframes(pcm.tobytes())
        exported.append(
            {
                "sample_id": source["sample_id"],
                "label": source["label"],
                "actor_id": source["actor_id"],
                "transcript": source["transcript"],
                "recognition_rate": source["recognition_rate"],
                "source_audio_path": source["audio_path"],
                "processed_audio_path": str(target.resolve()),
                "processing": f"float32 -> channel mean -> resample {target_sample_rate} Hz; no VAD",
            }
        )
    return exported


def create_loso_splits(
    rows: list[dict[str, Any]], output_dir: Path, validation_offset: int, seed: int
) -> list[dict[str, Any]]:
    actors = sorted({int(row["actor_id"]) for row in rows})
    reports: list[dict[str, Any]] = []
    fields = list(rows[0]) + ["split", "fold"]
    for test_index, test_actor in enumerate(actors):
        validation_actor = actors[(test_index + validation_offset) % len(actors)]
        fold_rows: dict[str, list[dict[str, Any]]] = {"train": [], "validation": [], "test": []}
        for source in rows:
            actor = int(source["actor_id"])
            split = "test" if actor == test_actor else "validation" if actor == validation_actor else "train"
            fold_rows[split].append({**source, "split": split, "fold": test_actor})
        for split, values in fold_rows.items():
            write_csv(output_dir / "splits" / f"fold_{test_actor:02d}_{split}.csv", values, fields)

        speakers = {
            split: sorted({int(row["actor_id"]) for row in values})
            for split, values in fold_rows.items()
        }
        texts = {split: {row["transcript"] for row in values} for split, values in fold_rows.items()}
        report = {
            "fold": test_actor,
            "seed": seed,
            "test_actor": test_actor,
            "validation_actor": validation_actor,
            "train_actors": speakers["train"],
            "samples": {split: len(values) for split, values in fold_rows.items()},
            "speaker_overlap_train_test": sorted(set(speakers["train"]) & set(speakers["test"])),
            "speaker_overlap_train_validation": sorted(
                set(speakers["train"]) & set(speakers["validation"])
            ),
            "text_overlap_unique_train_test": len(texts["train"] & texts["test"]),
            "text_overlap_unique_train_validation": len(texts["train"] & texts["validation"]),
            "label_counts": {
                split: dict(Counter(row["label"] for row in values))
                for split, values in fold_rows.items()
            },
        }
        if report["speaker_overlap_train_test"] or report["speaker_overlap_train_validation"]:
            raise ValueError(f"fold {test_actor} 存在演员泄漏")
        reports.append(report)
    return reports


def main() -> None:
    parser = argparse.ArgumentParser(description="MCAE-SPPS 数据、标签和 LOSO 划分审计")
    parser.add_argument("--config", required=True)
    parser.add_argument("--no-export-audio", action="store_true")
    args = parser.parse_args()
    config = load_config(args.config)
    paths = config["paths"]
    audit_config = config["audit"]
    output_dir = resolve_project_path(paths["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)

    metadata_path = resolve_project_path(paths["metadata_xlsx"])
    rater_path = resolve_project_path(paths["rater_xlsx"])
    manifest_path = resolve_project_path(paths["existing_manifest"])
    audio_dir = resolve_project_path(paths["audio_dir"])
    gender_by_actor = actor_gender_map(rater_path)
    rows, label_audit = build_master_rows(
        metadata_path, audio_dir, manifest_path, gender_by_actor
    )
    audio_rows, audio_audit = audit_audio(
        rows,
        float(audit_config["silence_threshold_dbfs"]),
        float(audit_config["clipping_threshold"]),
    )

    master_fields = list(rows[0])
    write_csv(output_dir / "master_manifest.csv", rows, master_fields)
    write_csv(output_dir / "audio_audit.csv", audio_rows, list(audio_rows[0]))

    seed = int(audit_config["seed"])
    random_generator = random.Random(seed)
    sample_count = int(audit_config["samples_per_class"])
    selected = [
        source
        for label in LABELS
        for source in random_generator.sample(
            [row for row in rows if row["label"] == label], sample_count
        )
    ]
    listening_fields = [
        "sample_id",
        "label",
        "actor_id",
        "transcript",
        "recognition_rate",
        "source_audio_path",
        "processed_audio_path",
        "processing",
    ]
    export_audio = bool(audit_config["export_processed_audio"]) and not args.no_export_audio
    if export_audio:
        listening_rows = export_listening_samples(
            selected, output_dir, int(audit_config["target_sample_rate"])
        )
    else:
        listening_rows = [
            {
                "sample_id": row["sample_id"],
                "label": row["label"],
                "actor_id": row["actor_id"],
                "transcript": row["transcript"],
                "recognition_rate": row["recognition_rate"],
                "source_audio_path": row["audio_path"],
                "processed_audio_path": "",
                "processing": "not exported",
            }
            for row in selected
        ]
    write_csv(output_dir / "listening_samples.csv", listening_rows, listening_fields)

    split_reports = create_loso_splits(
        rows,
        output_dir,
        int(config["split"]["validation_offset"]),
        seed,
    )
    full_audit = {
        "source_files": {
            "metadata_xlsx": str(metadata_path),
            "rater_xlsx": str(rater_path),
            "existing_manifest": str(manifest_path),
            "audio_dir": str(audio_dir),
        },
        "label_audit": {
            **label_audit,
            "label_counts": dict(Counter(row["label"] for row in rows)),
            "actor_counts": dict(Counter(str(row["actor_id"]) for row in rows)),
            "actor_gender": {str(key): value for key, value in gender_by_actor.items()},
        },
        "audio_audit": audio_audit,
        "inference_pipeline": {
            "required_input": "mono, 16000 Hz, float32",
            "funasr_reduce_channels": "mean over channel axis (reduce_channels=True)",
            "funasr_resample_target_hz": 16000,
            "emotion2vec_layer_norm": True,
            "vad_enabled": False,
            "model_output_order": MODEL_OUTPUT_ORDER,
        },
        "listening_sample_seed": seed,
        "listening_samples_per_class": sample_count,
        "processed_samples_exported": export_audio,
        "loso": split_reports,
    }
    write_json(output_dir / "audit_summary.json", full_audit)
    print(json.dumps(full_audit, ensure_ascii=False, indent=2))
    print(f"审计结果已保存: {output_dir}")


if __name__ == "__main__":
    main()
