"""Fresh downloads must not require an author's private manifest."""
import csv
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from prepare_mcae_manifest import prepare


class ManifestPreparationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.audio = self.root / "audio/Anger/123.wav"
        self.audio.parent.mkdir(parents=True)
        self.audio.write_bytes(b"synthetic fixture, not decoded by manifest generation")
        self.metadata = self.root / "metadata.xlsx"
        workbook = Workbook()
        workbook.active.append(["Filename", "emotion", "Actress"])
        workbook.active.append([123, "anger", 1])
        workbook.save(self.metadata)
        workbook.close()
        self.config = {"paths": {"metadata_xlsx": str(self.metadata),
                                 "audio_dir": str(self.root / "audio"),
                                 "existing_manifest": str(self.root / "manifest.csv")}}

    def test_builds_manifest_from_metadata_and_preserves_existing_by_default(self):
        path = prepare(self.config)
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(rows[0]["sample_id"], "123")
        self.assertEqual(rows[0]["label"], "angry")
        self.assertEqual(rows[0]["actor_id"], "1")
        self.assertEqual(Path(rows[0]["audio_path"]), self.audio.resolve())
        before = path.read_bytes()
        with self.assertRaises(FileExistsError):
            prepare(self.config)
        self.assertEqual(path.read_bytes(), before)

    def test_missing_audio_does_not_write_partial_manifest(self):
        self.audio.unlink()
        with self.assertRaises(FileNotFoundError):
            prepare(self.config)
        self.assertFalse(Path(self.config["paths"]["existing_manifest"]).exists())

    def test_duplicate_metadata_is_rejected(self):
        from openpyxl import load_workbook
        workbook = load_workbook(self.metadata)
        workbook.active.append([123, "anger", 2])
        workbook.save(self.metadata)
        workbook.close()
        with self.assertRaises(ValueError):
            prepare(self.config)


if __name__ == "__main__":
    unittest.main()
